# -*- coding: utf-8 -*-
"""Grava a aba 'Acionistas' no gnlink-bp-historico.xlsx: aporte de capital ACUMULADO por acionista
(Lorinvest/Hankoe e Copa Energia), cronograma histórico EXATO extraído da aba 'Exposição' do modelo
GNLink_Model_2026.09.07.xlsx (colunas D=Aportes Lorinvest, E=Aportes Copa, cumulativos em R$mi).
Copa entrou em ago/25 com lump de 100mi por 36%; Lorinvest aportou ~101mi ao longo de 2022-24.
Padrão igual ao patch_mincash.py (a série vem do 09.07, não do 09.04 que a base extrai). Regenera o mirror .js."""
import openpyxl, base64, re, html, zipfile, datetime, os
BP=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XL=os.path.join(BP,"gnlink-bp-historico.xlsx"); JS=os.path.join(BP,"gnlink-bp-historico.js")
M=r"C:/Users/murilo.nunes/OneDrive - Lorinvest Gestão de Recursos LTDA/Investimentos - Empresas/Lorinvest/Plano de Negócios/lorinvest-research/empresas/gnlink/analises-financeiras/Modelo/GNLink_Model_2026.09.07.xlsx"
Z=zipfile.ZipFile(M); wbx=Z.read('xl/workbook.xml').decode('utf8'); rels=Z.read('xl/_rels/workbook.xml.rels').decode('utf8')
sheets=re.findall(r'<sheet[^>]*name="([^"]+)"[^>]*r:id="([^"]+)"',wbx); relmap=dict(re.findall(r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"',rels))
n2f={html.unescape(n):relmap[r] for n,r in sheets}; en=[k for k in n2f if k.startswith("Exposi")][0]
xml=Z.read('xl/'+n2f[en]).decode('utf8')
def gv(col,r):
    m=re.search(r'<c r="'+col+str(r)+r'"([^>]*?)(?:/>|>(.*?)</c>)',xml,re.S)
    if not m: return None
    v=re.search(r'<v>(.*?)</v>',m.group(2) or ''); return float(v.group(1)) if v else None
def d2ym(s):
    dt=datetime.date(1899,12,30)+datetime.timedelta(days=int(round(s))); return "%04d-%02d"%(dt.year,dt.month)
hist={}
for r in range(16,200):
    c=gv("C",r)
    if c is None: break
    hist[d2ym(c)]=(gv("D",r) or 0.0, gv("E",r) or 0.0)
firstym=min(hist)
wb=openpyxl.load_workbook(XL); bl=wb["Balanco"]; yms=[bl.cell(1,c).value for c in range(2,bl.max_column+1)]
lor=[]; cop=[]; lastL=0.0; lastC=0.0
for ym in yms:
    if ym in hist: lastL,lastC=hist[ym]
    elif ym<firstym: lastL,lastC=0.0,0.0
    lor.append(lastL); cop.append(lastC)
if "Acionistas" in wb.sheetnames: del wb["Acionistas"]
ws=wb.create_sheet("Acionistas"); ws.append(["field"]+list(yms))
ws.append(["lorinvest_cum"]+lor); ws.append(["copa_cum"]+cop)
wb.save(XL)
open(JS,"w",encoding="ascii").write('window.__BP_XLSX_B64="'+base64.b64encode(open(XL,"rb").read()).decode("ascii")+'";\n')
print("Acionistas gravada:", firstym, "->", max(hist), "| Lorinvest final", lor[-1], "Copa final", cop[-1])
