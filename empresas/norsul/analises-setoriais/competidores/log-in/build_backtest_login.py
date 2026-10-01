"""Gera Log-In_Backtest_Volume.xlsx: modelo de volume (capacidade -> TEU) da Log-In, com backtest.

Logica: TEU(servico, trade, tri) = Capacidade nominal x dias operando / rotacao (dias)  [= slot-viagens ofertadas]
                                    x produtividade (TEU por slot-viagem) x mix de trade do servico.
Fatos publicados (releases 1T22-2T26, planilha de RI, programacao de navios set/26) entram como premissas verdes;
a serie trimestral por trade e reconstruida por formula. Os coeficientes sao calibrados aqui (NNLS) e gravados
como premissas; o Excel recalcula todo o resto.

Uso: python build_backtest_login.py   (gera o .xlsx na mesma pasta; recalcular no Excel ao abrir)
"""
import datetime as dt
import itertools
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("Log-In_Backtest_Volume.xlsx")
D = dt.date

# ----------------------------------------------------------------------------------------------------------
# 1. DADOS
# ----------------------------------------------------------------------------------------------------------
QL = ["1T22", "2T22", "3T22", "4T22", "1T23", "2T23", "3T23", "4T23", "1T24", "2T24", "3T24", "4T24",
      "1T25", "2T25", "3T25", "4T25", "1T26", "2T26"]
NQ = len(QL)
QS = [D(2022 + i // 4, 3 * (i % 4) + 1, 1) for i in range(NQ)]
QE = [(D(s.year + (s.month == 10), (s.month + 3 - 1) % 12 + 1, 1) - dt.timedelta(days=1)) for s in QS]

# Total navegacao (mil TEU) - Planilha de Resultados de RI (aba Volumes)
TOTAL = [98.9, 116.3, 114.1, 112.7, 102.0, 101.4, 136.0, 132.8, 155.7, 205.0, 198.6, 173.3,
         194.1, 181.3, 201.4, 200.1, 184.7, 193.4]

# Fatos publicados por trade: (id, trade, descricao, valor, tipo, fonte)
FACTS = [
    ("C22A", "Cabotagem", "Volume 2022 (ano)", 157.6, "v", "Release 4T22 - destaques"),
    ("C3T22", "Cabotagem", "Volume 3T22", 41.8, "v", "Release 3T22 - destaques"),
    ("C2T23", "Cabotagem", "Volume 2T23", 46.5, "v", "Release 2T23 - Volumes"),
    ("g_C2T23", "Cabotagem", "Variação 2T23 vs 2T22", 0.13, "g", "Release 2T23 - Volumes"),
    ("C3T23", "Cabotagem", "Volume 3T23", 54.5, "v", "Release 3T23 - Volumes"),
    ("C4T23", "Cabotagem", "Volume 4T23", 40.6, "v", "Release 4T23 - Volumes"),
    ("C23A", "Cabotagem", "Volume 2023 (ano)", 182.5, "v", "Release 4T23 - destaques"),
    ("g_C1T23", "Cabotagem", "Variação 1T23 vs 1T22", 0.19, "g", "Release 1T23 - Volumes"),
    ("C1T24", "Cabotagem", "Volume 1T24", 52.6, "v", "Release 1T24 - Volumes"),
    ("C2T24", "Cabotagem", "Volume 2T24", 56.8, "v", "Release 2T24 - Volumes"),
    ("g_C3T24", "Cabotagem", "Variação 3T24 vs 3T23", -0.04, "g", "Release 3T24 - Volumes"),
    ("g_C4T24", "Cabotagem", "Variação 4T24 vs 4T23", 0.016, "g", "Release 4T24 - Volumes"),
    ("g_C24A", "Cabotagem", "Variação 2024 vs 2023", 0.112, "g", "Release 4T24 - Volumes (checagem)"),
    ("g_C1T25", "Cabotagem", "Variação 1T25 vs 1T24", -0.176, "g", "Release 1T25 - Volumes"),
    ("C2T25", "Cabotagem", "Volume 2T25", 57.1, "v", "Release 2T25 - Volumes"),
    ("C3T25", "Cabotagem", "Volume 3T25", 71.6, "v", "Release 3T25 - Volumes"),
    ("C4T25", "Cabotagem", "Volume 4T25", 61.8, "v", "Release 4T25 - Volumes"),
    ("C25A", "Cabotagem", "Volume 2025 (ano)", 233.7, "v", "Release 4T25 - Volumes (checagem)"),
    ("F3T23", "Feeder", "Volume 3T23", 71.1, "v", "Release 3T23 - Volumes"),
    ("F4T23", "Feeder", "Volume 4T23", 82.9, "v", "Release 4T23 - Volumes"),
    ("F1T24", "Feeder", "Volume 1T24", 94.7, "v", "Release 1T24 - Volumes"),
    ("g_F1T24", "Feeder", "Variação 1T24 vs 1T23", 0.98, "g", "Release 1T24 - Volumes"),
    ("F2T24", "Feeder", "Volume 2T24", 138.0, "v", "Release 2T24 - Volumes"),
    ("g_F2T24", "Feeder", "Variação 2T24 vs 2T23", 2.12, "g", "Release 2T24 - Volumes"),
    ("F3T24", "Feeder", "Volume 3T24", 134.1, "v", "Release 3T24 - Volumes"),
    ("g_F3T24", "Feeder", "Variação 3T24 vs 3T23", 0.89, "g", "Release 3T24 - Volumes (checagem)"),
    ("F4T24", "Feeder", "Volume 4T24", 121.8, "v", "Release 4T24 - Volumes"),
    ("F24A", "Feeder", "Volume 2024 (ano)", 488.6, "v", "Release 4T24 - Volumes (checagem)"),
    ("F1T25", "Feeder", "Volume 1T25", 140.9, "v", "Release 1T25 - Volumes"),
    ("g_F2T25", "Feeder", "Variação 2T25 vs 2T24", -0.18, "g", "Release 2T25 - Volumes"),
    ("g_F3T25", "Feeder", "Variação 3T25 vs 3T24", -0.127, "g", "Release 3T25 - Volumes"),
    ("g_F4T25", "Feeder", "Variação 4T25 vs 4T24", 0.041, "g", "Release 4T25 - Volumes"),
    ("F25A", "Feeder", "Volume 2025 (ano)", 497.8, "v", "Release 4T25 - Volumes (checagem)"),
    ("M1T22", "Mercosul", "Volume 1T22", 12.6, "v", "Release 1T22 - destaques"),
    ("M2T22", "Mercosul", "Volume 2T22", 13.9, "v", "Release 2T22 - destaques"),
    ("g_M1T22", "Mercosul", "Variação 1T22 vs 1T21", 0.88, "g", "Release 1T22 - Volumes"),
    ("g_M2T22", "Mercosul", "Variação 2T22 vs 2T21", 0.52, "g", "Release 2T22 - Volumes"),
    ("g_M3T22", "Mercosul", "Variação 3T22 vs 3T21", 0.25, "g", "Release 3T22 - Volumes"),
    ("g_M9M22", "Mercosul", "Variação 9M22 vs 9M21", 0.50, "g", "Release 3T22 - Volumes"),
    ("g_M4T22", "Mercosul", "Variação 4T22 vs 4T21", 0.15, "g", "Release 4T22 - Volumes"),
    ("g_M22A", "Mercosul", "Variação 2022 vs 2021", 0.40, "g", "Release 4T22 - Volumes"),
    ("g_M1T23", "Mercosul", "Variação 1T23 vs 1T22", 0.06, "g", "Release 1T23 - Volumes (checagem)"),
    ("g_M3T24", "Mercosul", "Variação 3T24 vs 3T23", 0.18, "g", "Release 3T24 - Volumes (checagem)"),
    ("g_M4T24", "Mercosul", "Variação 4T24 vs 4T23", 0.098, "g", "Release 4T24 - Volumes (checagem)"),
    ("g_M24A", "Mercosul", "Variação 2024 vs 2023", -0.059, "g", "Release 4T24 - Volumes (checagem)"),
    ("g_M1T25", "Mercosul", "Variação 1T25 vs 1T24", 0.17, "g", "Release 1T25 - Volumes (checagem)"),
    ("g_M2T25", "Mercosul", "Variação 2T25 vs 2T24", 0.081, "g", "Release 2T25 - Volumes (checagem)"),
    ("g_M3T25", "Mercosul", "Variação 3T25 vs 3T24", 0.044, "g", "Release 3T25 - Volumes (checagem)"),
    ("g_M4T25", "Mercosul", "Variação 4T25 vs 4T24", 0.137, "g", "Release 4T25 - Volumes (checagem)"),
    ("g_M25A", "Mercosul", "Variação 2025 vs 2024", 0.102, "g", "Release 4T25 - Volumes (checagem)"),
    ("T22A", "Total", "Volume 2022 (ano)", 442.0, "v", "Release 4T22 (checagem)"),
    ("T23A", "Total", "Volume 2023 (ano)", 472.1, "v", "Release 4T23 (checagem)"),
    ("T24A", "Total", "Volume 2024 (ano)", 732.6, "v", "Release 4T24 (checagem)"),
    ("T25A", "Total", "Volume 2025 (ano)", 776.9, "v", "Release 4T25 (checagem)"),
]
FV = {f[0]: f[3] for f in FACTS}

# Frota: navio -> (modalidade, capacidade nominal TEU, observacao)
FLEET = [
    ("Log-In Polaris", "Próprio", 2700, "Na frota em todo o período"),
    ("Log-In Jacarandá", "Próprio", 2800, "Na frota em todo o período"),
    ("Log-In Jatobá", "Próprio", 2800, "Na frota em todo o período"),
    ("Log-In Endurance", "Próprio", 2800, "Na frota em todo o período"),
    ("Log-In Resiliente", "Próprio", 2700, "Vendido à MSC e afretado de volta (comunicado de 03/06/2026)"),
    ("Log-In Pantanal", "Próprio", 1700, "Vendido à MSC e afretado de volta (comunicado de 29/04/2026)"),
    ("Log-In Discovery", "Próprio", 2550, "Comprado em fev/21; fretado a terceiros (internacional) até mai/23"),
    ("MSC Belmonte III", "Afretado", 3500, "Afretado de jun/23 a jul/24 (3.534 TEU na tabela do 2T24)"),
    ("Log-In Evolution", "Próprio", 3158, "Entregue e em operação no 1T24 (3.150 TEU na tabela do 1T24)"),
    ("Log-In Experience", "Próprio", 3158, "Chegou ao Brasil em 23/07/2024; opera a partir do 3T24"),
]
CAP = {n: c for n, _, c, _ in FLEET}

# Servicos: codigo, nome, rotacao (dias), giro ilustrativo, rota, frequencia/observacao, fonte da rotacao
SERV = [
    ("SAM", "Serviço Amazonas", 35, 3.0, "Sul ↔ Norte (Manaus); encerrado no 3T23",
     "1 navio Log-In (Polaris) em rota com parceiros", "Premissa: igual à do SAS (sem programação histórica)"),
    ("SAS", "Serviço Atlântico Sul", 35, 3.0,
     "BUE-LPG-RIG-NVT-SSZ-SUA-PEC-SSA-SSZ-NVT-BUE",
     "Semanal; 5 posições, incluindo navio parceiro (M. Suape)",
     "Programação de navios 30/09/26: Resiliente v.498→499 em BUE com 35 dias"),
    ("SEA", "Serviço Expresso Amazonas (até abr/25)", 28, 3.0, "SSZ-NVT-SUA-PEC-MAO-SUA-VIX-SSZ",
     "Quinzenal (mai/23) → semanal (set/23); 3 navios Log-In + 1 de parceiro (troca de espaço)",
     "Release 2T23/3T23; volta de 28 dias = 4 posições × 7 dias"),
    ("SEA25", "Serviço Expresso Amazonas reformulado (desde abr/25)", 28, 3.5,
     "SSZ-NVT-SSA-SUA-PEC-ITA-MAO-SUA-VIX-SSZ",
     "Semanal; 4 navios Log-In; escala adicional em Salvador",
     "Programação 30/09/26: Jacarandá v.197→198 em 28 dias; Endurance e Evolution em 27 dias"),
    ("SSV", "Shuttle Vitória", 7, 2.0, "RIO-VIX-RIO", "Semanal; 1 navio",
     "Programação 30/09/26: Discovery v.80-87 em escalas semanais"),
    ("SSR", "Shuttle Rio", 7, 2.0, "SSZ-RIO-VIX-SSZ", "Semanal; 1 navio",
     "Programação 30/09/26: Pantanal v.614-621 em escalas semanais"),
    ("SSN", "Shuttle Navegantes (mai/24 a abr/25)", 7, 2.0, "Navegantes ↔ portos de transbordo",
     "Semanal (premissa); 1 navio (Evolution)", "Premissa: igual aos demais shuttles"),
]
SCODES = [s[0] for s in SERV]
OFFHIRE = [("DOC", "Docagem (fora de operação)"), ("FRET", "Fretado a terceiros (fora da rede Log-In)")]

END = D(2026, 6, 30)
# Alocacao: navio, servico, inicio, fim, fundamento
STINTS = [
    ("Log-In Polaris", "SAM", D(2022, 1, 1), D(2023, 8, 31), "Tabelas de frota 1T22-3T23: SAM"),
    ("Log-In Polaris", "SEA", D(2023, 9, 1), D(2025, 4, 20), "SEA semanal desde set/23 com Polaris, Jatobá e Discovery (Release 2T23)"),
    ("Log-In Polaris", "SEA25", D(2025, 4, 21), END, "Tabelas 2T25-4T25: SEA; 2026 sem tabela publicada (mantida)"),
    ("Log-In Jacarandá", "SAS", D(2022, 1, 1), D(2022, 8, 14), "Tabelas de frota 2022: SAS"),
    ("Log-In Jacarandá", "DOC", D(2022, 8, 15), D(2022, 10, 5), "Docagem 3T22, concluída no 4T22 (datas estimadas)"),
    ("Log-In Jacarandá", "SAS", D(2022, 10, 6), D(2023, 12, 31), "Tabelas 2023: SAS"),
    ("Log-In Jacarandá", "SEA", D(2024, 1, 1), D(2025, 4, 20), "Tabela 1T24: SEA"),
    ("Log-In Jacarandá", "SEA25", D(2025, 4, 21), END, "Tabelas 2T25-4T25: SEA"),
    ("Log-In Jatobá", "SAS", D(2022, 1, 1), D(2023, 8, 31), "Tabelas 2022-3T23: SAS"),
    ("Log-In Jatobá", "SEA", D(2023, 9, 1), D(2024, 12, 31), "Tabelas 4T23-4T24: SEA"),
    ("Log-In Jatobá", "SAS", D(2025, 1, 1), END, "Tabelas 1T25-4T25: SAS"),
    ("Log-In Endurance", "SAS", D(2022, 1, 1), D(2026, 5, 15), "Tabelas 2022-4T25: SAS"),
    ("Log-In Endurance", "DOC", D(2026, 5, 16), D(2026, 5, 31), "Docagem 2T26 (Release 2T26); 16 dias calibrados p/ capacidade média de 23.389"),
    ("Log-In Endurance", "SAS", D(2026, 6, 1), END, "Volta ao serviço após docagem"),
    ("Log-In Resiliente", "SSV", D(2022, 1, 1), D(2026, 6, 14), "Serviço Shuttle; SSV explícito desde 2T24"),
    ("Log-In Resiliente", "DOC", D(2026, 6, 15), END, "Docagem 2T26 (Release 2T26); 16 dias calibrados"),
    ("Log-In Pantanal", "SSR", D(2022, 1, 1), D(2022, 9, 30), "Serviço Shuttle; SSR explícito desde 2T24"),
    ("Log-In Pantanal", "DOC", D(2022, 10, 1), D(2022, 11, 15), "Docagem 3T22/4T22 restringiu o Feeder no 4T22 (datas estimadas)"),
    ("Log-In Pantanal", "SSR", D(2022, 11, 16), END, "Serviço Shuttle / SSR"),
    ("Log-In Discovery", "FRET", D(2022, 1, 1), D(2023, 6, 9), "Fretamento internacional em 2022 (Releases 1T22/2T22)"),
    ("Log-In Discovery", "SEA", D(2023, 6, 10), D(2024, 2, 14), "Entra na cabotagem em jun/23 (Release 3T23)"),
    ("Log-In Discovery", "DOC", D(2024, 2, 15), D(2024, 4, 15), "Em docagem no fim do 1T24 (tabela 1T24); datas estimadas"),
    ("Log-In Discovery", "SAS", D(2024, 4, 16), END, "Tabelas 2T24-4T25: SAS"),
    ("MSC Belmonte III", "SEA", D(2023, 6, 10), D(2023, 8, 31), "Afretado em jun/23 para o SEA (Release 3T23)"),
    ("MSC Belmonte III", "SAS", D(2023, 9, 1), D(2024, 7, 15), "Tabelas 4T23-2T24: SAS; devolvido em jul/24"),
    ("Log-In Evolution", "SAS", D(2024, 2, 15), D(2024, 4, 30), "Entregue no 1T24; tabela 1T24: SAS"),
    ("Log-In Evolution", "SSN", D(2024, 5, 1), D(2025, 4, 20), "SSN lançado em mai/24 com o Evolution; encerrado em abr/25"),
    ("Log-In Evolution", "SEA25", D(2025, 4, 21), END, "Realocado ao SEA após o fim do SSN (Release 4T25)"),
    ("Log-In Experience", "SAS", D(2024, 8, 15), D(2024, 12, 31), "Início de operação no SAS no 3T24 (data estimada)"),
    ("Log-In Experience", "SEA", D(2025, 1, 1), D(2025, 4, 20), "Tabela 1T25: SEA"),
    ("Log-In Experience", "SEA25", D(2025, 4, 21), END, "Tabelas 2T25-4T25: SEA"),
]

# Capacidade nominal reportada (TEU) - linha "Frota - Capacidade Nominal" dos releases
CAP_REP = [18050, 18050, 18050, 18050, 18050, 21550, 21550, 21550, 22150, 24742, 24366, 24366,
           24366, 24366, 24366, 24366, 24366, 23389]

EVENTS = [
    "Discovery fretado fora da rede; Mercosul recorde p/ 1T",
    "Mercosul recorde p/ 2T (novos clientes)",
    "Docagens de Jacarandá e Pantanal; Feeder em take-or-pay (receita sem volume)",
    "Pantanal em docagem (Feeder restrito); seca no Amazonas",
    "Feeder fraco (queda de armadores); licenças de importação na Argentina",
    "Crise argentina; fretes internacionais em queda; SEA começa quinzenal (mai/23)",
    "SEA semanal (set/23); Feeder contingencial (mau tempo no Sul); início da seca",
    "Seca no Amazonas (balsas) e cheia em Itajaí (omissões em Navegantes)",
    "Discovery em docagem; Feeder contingencial recorde; Evolution entra",
    "SSN lançado (mai/24) + portos congestionados: embarques contingenciais",
    "Belmonte devolvido; Experience entra; omissões de escala no SAS/SEA",
    "Congestionamento + seca (píer de Itacoatiara): Cabotagem penalizada",
    "Novo player na cabotagem; Feeder de-para Manaus; chuvas em Bahía Blanca",
    "Fim do SSN (abr/25); SEA reformulado c/ 4 navios e Salvador",
    "Cabotagem recorde (+37%): recuperação do nível de serviço",
    "Cabotagem +50% sobre base fraca (seca do 4T24)",
    "Feeder cai pós-SSN; Cabotagem recorde p/ 1T; exportações argentinas",
    "Docagens Endurance/Resiliente; frete mínimo ANTT favorece cabotagem",
]


# ----------------------------------------------------------------------------------------------------------
# 2. SERIE RECONSTRUIDA + CALIBRACAO (espelha as formulas do Excel)
# ----------------------------------------------------------------------------------------------------------
def series():
    f = FV
    cab = [None] * NQ
    cab[2], cab[5], cab[6], cab[7] = f["C3T22"], f["C2T23"], f["C3T23"], f["C4T23"]
    cab[4] = f["C23A"] - f["C2T23"] - f["C3T23"] - f["C4T23"]
    cab[0] = cab[4] / (1 + f["g_C1T23"])
    cab[1] = f["C2T23"] / (1 + f["g_C2T23"])
    cab[3] = f["C22A"] - cab[0] - cab[1] - cab[2]
    cab[8], cab[9] = f["C1T24"], f["C2T24"]
    cab[10] = cab[6] * (1 + f["g_C3T24"])
    cab[11] = cab[7] * (1 + f["g_C4T24"])
    cab[12] = cab[8] * (1 + f["g_C1T25"])
    cab[13], cab[14], cab[15] = f["C2T25"], f["C3T25"], f["C4T25"]
    fee = [None] * NQ
    fee[6], fee[7], fee[8], fee[9], fee[10], fee[11] = f["F3T23"], f["F4T23"], f["F1T24"], f["F2T24"], f["F3T24"], f["F4T24"]
    fee[4] = f["F1T24"] / (1 + f["g_F1T24"])
    fee[5] = f["F2T24"] / (1 + f["g_F2T24"])
    fee[12] = f["F1T25"]
    fee[13] = fee[9] * (1 + f["g_F2T25"])
    fee[14] = fee[10] * (1 + f["g_F3T25"])
    fee[15] = fee[11] * (1 + f["g_F4T25"])
    mer = [None] * NQ
    a = f["M1T22"] / (1 + f["g_M1T22"]); b = f["M2T22"] / (1 + f["g_M2T22"])
    c = (f["M1T22"] + f["M2T22"] - (1 + f["g_M9M22"]) * (a + b)) / ((1 + f["g_M9M22"]) - (1 + f["g_M3T22"]))
    mer[0], mer[1], mer[2] = f["M1T22"], f["M2T22"], c * (1 + f["g_M3T22"])
    d = (mer[0] + mer[1] + mer[2] - (1 + f["g_M22A"]) * (a + b + c)) / ((1 + f["g_M22A"]) - (1 + f["g_M4T22"]))
    mer[3] = d * (1 + f["g_M4T22"])
    for i in range(4):
        fee[i] = TOTAL[i] - cab[i] - mer[i]
    for i in range(4, 16):
        mer[i] = TOTAL[i] - cab[i] - fee[i]
    return cab, mer, fee


CAB, MER, FEE = series()
ROT = {s[0]: s[2] for s in SERV}


def offers():
    o = {s: np.zeros(NQ) for s in SCODES}
    for ship, s, a, b, _ in STINTS:
        if s not in ROT:
            continue
        for i in range(NQ):
            dd = (min(b, QE[i]) - max(a, QS[i])).days + 1
            if dd > 0:
                o[s][i] += CAP[ship] * dd / ROT[s] / 1000
    return o


OFF = offers()
SH = OFF["SSV"] + OFF["SSR"] + OFF["SSN"]


def nnls(A, y):
    best, bres = None, 1e99
    n = A.shape[1]
    for k in range(1, n + 1):
        for sub in itertools.combinations(range(n), k):
            c, *_ = np.linalg.lstsq(A[:, sub], y, rcond=None)
            if (c >= -1e-12).all():
                r = ((A[:, sub] @ c - y) ** 2).sum()
                if r < bres:
                    best = np.zeros(n); best[list(sub)] = c; bres = r
    return best


def calibrate(idx, tie_sea25):
    # theta = [a(shuttle), b1(feeder SAS), b2(cab SAS), b3(merc SAS), d1(feeder SEA), d2(cab SEA), e1, e2]
    rows, y = [], []
    for t in idx:
        s25 = OFF["SEA25"][t]
        seaX = OFF["SEA"][t] + (s25 if tie_sea25 else 0)
        rows.append([SH[t], OFF["SAS"][t], 0, 0, seaX, 0, 0 if tie_sea25 else s25, 0]); y.append(FEE[t])
        rows.append([0, OFF["SAM"][t], OFF["SAS"][t] + OFF["SAM"][t], OFF["SAM"][t], 0, seaX, 0, 0 if tie_sea25 else s25]); y.append(CAB[t])
        rows.append([0, 0, 0, OFF["SAS"][t], 0, 0, 0, 0]); y.append(MER[t])
    A = np.array(rows); y = np.array(y)
    keep = [j for j in range(A.shape[1]) if np.abs(A[:, j]).sum() > 0]
    th = np.zeros(A.shape[1]); th[keep] = nnls(A[:, keep], y)
    a, b1, b2, b3, d1, d2, e1, e2 = th
    if tie_sea25:
        e1, e2 = d1, d2
    pSAS = b1 + b2 + b3; pSEA = d1 + d2; p25 = e1 + e2
    # p e mix (cab, merc, feeder) por servico
    par = {
        "SAM": (pSAS, (1.0, 0.0, 0.0)),
        "SAS": (pSAS, (b2 / pSAS, b3 / pSAS, b1 / pSAS)),
        "SEA": (pSEA, (d2 / pSEA, 0.0, d1 / pSEA)),
        "SEA25": (p25, (e2 / p25, 0.0, e1 / p25)),
        "SSV": (a, (0.0, 0.0, 1.0)), "SSR": (a, (0.0, 0.0, 1.0)), "SSN": (a, (0.0, 0.0, 1.0)),
    }
    return par


PAR_A = calibrate(range(0, 12), True)   # 1T22-4T24; 2025-26 fora da amostra
PAR_B = calibrate(range(0, 16), False)  # 1T22-4T25


def predict(par):
    T = np.zeros(NQ)
    for s in SCODES:
        T += OFF[s] * par[s][0]
    return T


PRED_A, PRED_B = predict(PAR_A), predict(PAR_B)

# ----------------------------------------------------------------------------------------------------------
# 3. ESTILO
# ----------------------------------------------------------------------------------------------------------
ARIAL = "Arial"
F_N = Font(name=ARIAL, size=10)
F_B = Font(name=ARIAL, size=10, bold=True)
F_T = Font(name=ARIAL, size=13, bold=True, color="1F3864")
F_S = Font(name=ARIAL, size=10, bold=True, color="1F3864")
F_I = Font(name=ARIAL, size=9, italic=True, color="595959")
F_W = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
FILL_IN = PatternFill("solid", fgColor="E3ECE8")
FILL_H = PatternFill("solid", fgColor="1F3864")
FILL_SUB = PatternFill("solid", fgColor="F2F2F2")
FILL_OUT = PatternFill("solid", fgColor="FCE4D6")
BOT = Border(bottom=Side(style="thin", color="BFBFBF"))
TOP = Border(top=Side(style="thin", color="7F7F7F"))
NF1 = '_(#,##0.0_);_((#,##0.0);_("-"??_);_(@_)'
NF0 = '_(#,##0_);_((#,##0);_("-"??_);_(@_)'
NF2 = '_(#,##0.00_);_((#,##0.00);_("-"??_);_(@_)'
NFP = '0.0%;(0.0%);"-"'
NFD = "dd/mm/yy"
WRAP = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(horizontal="center")

wb = Workbook()
wb._named_styles["Normal"].font = Font(name=ARIAL, size=10)


def W(ws, ref, v, font=None, fill=None, nf=None, al=None, border=None):
    c = ws[ref]
    c.value = v
    c.font = font or F_N
    if fill: c.fill = fill
    if nf: c.number_format = nf
    if al: c.alignment = al
    if border: c.border = border
    return c


def title(ws, text, sub=None):
    W(ws, "A1", text, F_T)
    if sub: W(ws, "A2", sub, F_I)
    ws.sheet_view.showGridLines = True


def header_row(ws, row, c0, labels, fill=FILL_H, font=F_W):
    for j, lab in enumerate(labels):
        W(ws, f"{L(c0 + j)}{row}", lab, font, fill, al=Alignment(horizontal="center", vertical="center", wrap_text=True))
    if any(len(str(x)) > 11 for x in labels):
        ws.row_dimensions[row].height = 30


# ----------------------------------------------------------------------------------------------------------
# 4. ABAS
# ----------------------------------------------------------------------------------------------------------
# ---- Leia-me ----
ws = wb.active; ws.title = "Leia-me"
title(ws, "Log-In · Modelo de volume (capacidade → TEU) com backtest",
      "Navegação Costeira (Cabotagem, Mercosul, Feeder) · 1T22–2T26 · mil TEU · gerado por build_backtest_login.py")
lines = [
    ("Pergunta", "Quanto do volume publicado pela Log-In se explica pela capacidade instalada (navios × dias × rotação) e quanto é produtividade (ocupação/giro) que muda com o mercado?"),
    ("Fórmula", "TEU(serviço, trade, tri) = Capacidade nominal × Dias operando ÷ Rotação (dias) × Produtividade (TEU por slot-viagem) × Mix de trade do serviço"),
    ("", "Slot-viagens = capacidade ofertada no trimestre (cada viagem redonda oferece a capacidade nominal do navio uma vez)."),
    ("", "Produtividade = fator efetivo (peso a 14 t, reefer) × giro de slot (quantas vezes a vaga é revendida por volta) × ocupação. Só o produto é identificável pelos dados; a abertura em Servicos é ilustrativa."),
    ("Backtest", "Janela 1: calibra em 1T22–4T24 e prevê 1T25–2T26 fora da amostra (o SEA reformulado herda os parâmetros do SEA antigo). Janela 2: calibra em 1T22–4T25 (amostra completa). Trocar em Backtest!D3."),
    ("Reconciliação", "A aba Reconciliacao aplica o fator publicado ÷ modelado por trade e devolve a produtividade implícita de cada serviço e trimestre; com ela o modelo chega exatamente no volume publicado."),
    ("Abas", "Volumes_Publicados → fatos dos releases (verde) e série trimestral por trade reconstruída por fórmula"),
    ("", "Frota · Servicos · Alocacao → premissas (capacidades, rotações, navio × serviço × período, docagens)"),
    ("", "Calc_Navio → dias operando e slot-viagens por navio e trimestre · Calc_Servico → oferta e volume modelado por serviço × trade"),
    ("", "Backtest → publicado × modelado, erro, MAPE dentro/fora da amostra · Reconciliacao → produtividade implícita e ponte anual"),
    ("", "Checagens → capacidade de frota reportada × modelo e fechamento dos números anuais · Evidencia_Rotacao → programação de navios de set/26"),
    ("", "ANTAQ_Plug → layout para plugar a base por navio/escala da ANTAQ (painel fora do ar em out/26) · Fontes"),
    ("Cores", "Verde (#E3ECE8) = premissa/fato digitado · preto = fórmula · laranja = trimestre fora da amostra de calibração"),
    ("Limitações", "1) Troca de espaço com parceiros (SEA desde set/23 e SAS com o Mercosul Suape; Serviço Manaus e SPA em navios de terceiros) não é modelada: a capacidade é a dos navios operados pela Log-In."),
    ("", "2) Rotações de 2022–2025 assumidas iguais às observadas na programação de set/26 (Shuttle 7 d, SEA 28 d, SAS 35 d); SAM e SSN sem programação histórica."),
    ("", "3) Datas de docagem e de troca de serviço dentro do trimestre são estimadas a partir dos releases; o mapa navio × serviço vem das tabelas de frota trimestrais."),
    ("", "4) A abertura por trade de 2026 não é divulgada: em 2026 o backtest compara só o total."),
    ("", "5) Feeder 2022 e Mercosul 2023–2025 são resíduos (total publicado − demais trades); Mercosul 3T22/4T22 deriva das variações a/a de 9M e 12M."),
]
r = 4
for k, v in lines:
    W(ws, f"A{r}", k, F_B); W(ws, f"B{r}", v, al=WRAP); r += 1
ws.column_dimensions["A"].width = 16; ws.column_dimensions["B"].width = 150

# ---- Volumes_Publicados ----
wp = wb.create_sheet("Volumes_Publicados")
title(wp, "Volumes publicados e série trimestral por trade (mil TEU)",
      "Verde = fato publicado (release/planilha de RI). Preto = derivado por fórmula. Base indica como cada número foi obtido.")
W(wp, "A4", "Fatos publicados", F_S)
header_row(wp, 5, 1, ["ID", "Trade", "Descrição", "Valor", "Fonte"])
fact_row = {}
for i, (fid, tr, desc, val, typ, src) in enumerate(FACTS):
    rr = 6 + i
    W(wp, f"A{rr}", fid); W(wp, f"B{rr}", tr); W(wp, f"C{rr}", desc)
    W(wp, f"D{rr}", val, fill=FILL_IN, nf=(NFP if typ == "g" else NF1)); W(wp, f"E{rr}", src)
    fact_row[fid] = rr
FR = lambda fid: f"Volumes_Publicados!$D${fact_row[fid]}"
for col, w in zip("ABCDE", [10, 11, 30, 10, 38]):
    wp.column_dimensions[col].width = w

# serie trimestral: colunas H.. (H = 1T22)
C0 = 8
qcol = lambda i: L(C0 + i)
W(wp, "G4", "Série trimestral reconstruída", F_S)
W(wp, "G5", "mil TEU", F_B)
for i, q in enumerate(QL):
    W(wp, f"{qcol(i)}5", q, F_W, FILL_H, al=CEN)
RT, RC, RM, RF = 6, 7, 8, 9           # total, cab, merc, feeder
W(wp, f"G{RT}", "Total (planilha RI)", F_B)
W(wp, f"G{RC}", "Cabotagem"); W(wp, f"G{RM}", "Mercosul"); W(wp, f"G{RF}", "Feeder")
W(wp, "G10", "Soma dos trades − total", F_I)
for i in range(NQ):
    W(wp, f"{qcol(i)}{RT}", TOTAL[i], F_B, FILL_IN, NF1)
c = qcol
cab_f = {
    0: f"={c(4)}{RC}/(1+{FR('g_C1T23')})", 1: f"={FR('C2T23')}/(1+{FR('g_C2T23')})", 2: f"={FR('C3T22')}",
    3: f"={FR('C22A')}-{c(0)}{RC}-{c(1)}{RC}-{c(2)}{RC}", 4: f"={FR('C23A')}-{c(5)}{RC}-{c(6)}{RC}-{c(7)}{RC}",
    5: f"={FR('C2T23')}", 6: f"={FR('C3T23')}", 7: f"={FR('C4T23')}", 8: f"={FR('C1T24')}", 9: f"={FR('C2T24')}",
    10: f"={c(6)}{RC}*(1+{FR('g_C3T24')})", 11: f"={c(7)}{RC}*(1+{FR('g_C4T24')})", 12: f"={c(8)}{RC}*(1+{FR('g_C1T25')})",
    13: f"={FR('C2T25')}", 14: f"={FR('C3T25')}", 15: f"={FR('C4T25')}",
}
fee_f = {
    4: f"={FR('F1T24')}/(1+{FR('g_F1T24')})", 5: f"={FR('F2T24')}/(1+{FR('g_F2T24')})", 6: f"={FR('F3T23')}",
    7: f"={FR('F4T23')}", 8: f"={FR('F1T24')}", 9: f"={FR('F2T24')}", 10: f"={FR('F3T24')}", 11: f"={FR('F4T24')}",
    12: f"={FR('F1T25')}", 13: f"={c(9)}{RF}*(1+{FR('g_F2T25')})", 14: f"={c(10)}{RF}*(1+{FR('g_F3T25')})",
    15: f"={c(11)}{RF}*(1+{FR('g_F4T25')})",
}
# Mercosul 2021 implicito (auxiliar) em linhas 14-17
W(wp, "G13", "Mercosul 2021 implícito (auxiliar p/ 3T22 e 4T22)", F_S)
W(wp, "G14", "1T21 = 1T22 ÷ (1+var.)"); W(wp, "H14", f"={FR('M1T22')}/(1+{FR('g_M1T22')})", nf=NF1)
W(wp, "G15", "2T21 = 2T22 ÷ (1+var.)"); W(wp, "H15", f"={FR('M2T22')}/(1+{FR('g_M2T22')})", nf=NF1)
W(wp, "G16", "3T21: resolve var. 9M22 (+50%)")
W(wp, "H16", f"=({FR('M1T22')}+{FR('M2T22')}-(1+{FR('g_M9M22')})*(H14+H15))/((1+{FR('g_M9M22')})-(1+{FR('g_M3T22')}))", nf=NF1)
W(wp, "G17", "4T21: resolve var. 2022 (+40%)")
W(wp, "H17", f"=({c(0)}{RM}+{c(1)}{RM}+{c(2)}{RM}-(1+{FR('g_M22A')})*(H14+H15+H16))/((1+{FR('g_M22A')})-(1+{FR('g_M4T22')}))", nf=NF1)
mer_f = {0: f"={FR('M1T22')}", 1: f"={FR('M2T22')}", 2: f"=H16*(1+{FR('g_M3T22')})", 3: f"=H17*(1+{FR('g_M4T22')})"}
for i in range(NQ):
    if i in cab_f: W(wp, f"{c(i)}{RC}", cab_f[i], nf=NF1)
    else: W(wp, f"{c(i)}{RC}", "n.d.", al=CEN)
    if i < 4:
        W(wp, f"{c(i)}{RF}", f"={c(i)}{RT}-{c(i)}{RC}-{c(i)}{RM}", nf=NF1)
        W(wp, f"{c(i)}{RM}", mer_f[i], nf=NF1)
    elif i < 16:
        W(wp, f"{c(i)}{RF}", fee_f[i], nf=NF1)
        W(wp, f"{c(i)}{RM}", f"={c(i)}{RT}-{c(i)}{RC}-{c(i)}{RF}", nf=NF1)
    else:
        W(wp, f"{c(i)}{RF}", "n.d.", al=CEN); W(wp, f"{c(i)}{RM}", "n.d.", al=CEN)
    W(wp, f"{c(i)}10", f'=IF(ISNUMBER({c(i)}{RC}),{c(i)}{RC}+{c(i)}{RM}+{c(i)}{RF}-{c(i)}{RT},"")', F_I, nf=NF1)
# base (como cada numero foi obtido)
BASE_C = ["1T23 ÷ 1,19", "2T23 ÷ 1,13", "publicado", "2022 − 1T a 3T", "2023 − 2T a 4T", "publicado", "publicado",
          "publicado", "publicado", "publicado", "3T23 × 0,96", "4T23 × 1,016", "1T24 × 0,824", "publicado", "publicado",
          "publicado", "não divulgado", "não divulgado"]
BASE_F = ["resíduo", "resíduo", "resíduo", "resíduo", "1T24 ÷ 1,98", "2T24 ÷ 3,12", "publicado", "publicado", "publicado",
          "publicado", "publicado", "publicado", "publicado", "2T24 × 0,82", "3T24 × 0,873", "4T24 × 1,041",
          "não divulgado", "não divulgado"]
BASE_M = ["publicado", "publicado", "var. a/a 3T e 9M", "var. a/a 4T e 12M"] + ["resíduo"] * 12 + ["não divulgado"] * 2
W(wp, "G19", "Base do número", F_S)
for lab, arr, rr in [("Cabotagem", BASE_C, 20), ("Mercosul", BASE_M, 21), ("Feeder", BASE_F, 22)]:
    W(wp, f"G{rr}", lab)
    for i in range(NQ):
        W(wp, f"{c(i)}{rr}", arr[i], F_I, al=CEN)
# anuais
W(wp, "G25", "Anual (mil TEU)", F_S)
YRS = [("2022", 0), ("2023", 4), ("2024", 8), ("2025", 12), ("1S26", 16)]
for j, (y, s0) in enumerate(YRS):
    W(wp, f"{L(8 + j)}26", y, F_W, FILL_H, al=CEN)
for rr, lab, src in [(27, "Total", RT), (28, "Cabotagem", RC), (29, "Mercosul", RM), (30, "Feeder", RF)]:
    W(wp, f"G{rr}", lab, F_B if lab == "Total" else F_N)
    for j, (y, s0) in enumerate(YRS):
        n = 2 if y == "1S26" else 4
        rng = f"{c(s0)}{src}:{c(s0 + n - 1)}{src}"
        W(wp, f"{L(8 + j)}{rr}", f'=IF(COUNT({rng})={n},SUM({rng}),"n.d.")', nf=NF1, al=None)
wp.column_dimensions["G"].width = 34
for i in range(NQ):
    wp.column_dimensions[c(i)].width = 9.5
wp.freeze_panes = "H6"

# ---- Frota ----
wf = wb.create_sheet("Frota")
title(wf, "Frota operada pela Log-In (porta-contêineres)", "Capacidade nominal das tabelas de frota dos releases trimestrais")
header_row(wf, 4, 1, ["Navio", "Modalidade", "Capacidade nominal (TEU)", "IMO (p/ ANTAQ)", "Observação"])
for i, (n, m, cap, obs) in enumerate(FLEET):
    rr = 5 + i
    W(wf, f"A{rr}", n); W(wf, f"B{rr}", m); W(wf, f"C{rr}", cap, fill=FILL_IN, nf=NF0)
    W(wf, f"D{rr}", None, fill=FILL_IN); W(wf, f"E{rr}", obs)
FLEET_LAST = 4 + len(FLEET)
W(wf, f"A{FLEET_LAST + 2}", "IMO em branco: preencher para cruzar com a base de atracações da ANTAQ (aba ANTAQ_Plug).", F_I)
for col, w in zip("ABCDE", [22, 12, 22, 16, 70]):
    wf.column_dimensions[col].width = w

# ---- Servicos ----
wsv = wb.create_sheet("Servicos")
title(wsv, "Serviços: rotação, produtividade calibrada e mix de trade",
      "p = TEU transportado por slot-viagem ofertado. Calibrado por mínimos quadrados não-negativos (build_backtest_login.py). Janela selecionada em Backtest!D3.")
W(wsv, "A3", "Fator efetivo (TEU a 14 t ÷ nominal)", F_B); W(wsv, "D3", 0.85, fill=FILL_IN, nf=NF2)
W(wsv, "E3", "ilustrativo: só entra na decomposição da ocupação", F_I)
hdr = ["Código", "Serviço", "Rota (portos)", "Rotação (dias)", "Navios / frequência", "Giro de slot (ilustr.)",
       "p · janela 1", "p · janela 2", "p selecionado",
       "Cab · j1", "Merc · j1", "Feeder · j1", "Cab · j2", "Merc · j2", "Feeder · j2",
       "Cab · sel", "Merc · sel", "Feeder · sel", "Ocupação implícita", "Fonte da rotação"]
header_row(wsv, 5, 1, hdr)
SV_ROW = {}
for i, (code, name, rot, giro, rota, freq, src) in enumerate(SERV):
    rr = 6 + i; SV_ROW[code] = rr
    W(wsv, f"A{rr}", code, F_B); W(wsv, f"B{rr}", name); W(wsv, f"C{rr}", rota)
    W(wsv, f"D{rr}", rot, fill=FILL_IN, nf=NF0); W(wsv, f"E{rr}", freq); W(wsv, f"F{rr}", giro, fill=FILL_IN, nf=NF1)
    W(wsv, f"G{rr}", round(float(PAR_A[code][0]), 4), fill=FILL_IN, nf=NF2)
    W(wsv, f"H{rr}", round(float(PAR_B[code][0]), 4), fill=FILL_IN, nf=NF2)
    W(wsv, f"I{rr}", f"=CHOOSE(Janela,G{rr},H{rr})", F_B, nf=NF2)
    for k in range(3):
        W(wsv, f"{L(10 + k)}{rr}", round(float(PAR_A[code][1][k]), 4), fill=FILL_IN, nf=NFP)
        W(wsv, f"{L(13 + k)}{rr}", round(float(PAR_B[code][1][k]), 4), fill=FILL_IN, nf=NFP)
        W(wsv, f"{L(16 + k)}{rr}", f"=CHOOSE(Janela,{L(10 + k)}{rr},{L(13 + k)}{rr})", nf=NFP)
    W(wsv, f"S{rr}", f"=IFERROR(I{rr}/($D$3*F{rr}),0)", nf=NFP)
    W(wsv, f"T{rr}", src, F_I)
SV_FIRST, SV_LAST = 6, 6 + len(SERV) - 1
for j, (code, name) in enumerate(OFFHIRE):
    rr = SV_LAST + 1 + j
    W(wsv, f"A{rr}", code, F_B); W(wsv, f"B{rr}", name); W(wsv, f"D{rr}", 0, fill=FILL_IN, nf=NF0)
    W(wsv, f"T{rr}", "Sem oferta de slots (rotação 0)", F_I)
W(wsv, f"A{SV_LAST + 4}", "Leitura dos parâmetros", F_S)
notes = [
    "Shuttles (SSV/SSR/SSN): 100% Feeder; p ≈ 1 TEU por slot-viagem em voltas semanais de 2–4 trechos.",
    "SAS e SAM: mesma produtividade (restrição de calibração, pois 2022–23 não separa os dois); SAM 100% Cabotagem; o SAS divide Cabotagem e Mercosul.",
    "SEA (2023–abr/25): a calibração atribui a ele parte do Feeder contingencial de 2023–24 (cargas de-para Manaus e transbordo).",
    "SEA25 (desde abr/25): na janela 1 herda os parâmetros do SEA antigo (teste fora da amostra); na janela 2 é calibrado com 2T25–4T25.",
    "Ocupação implícita = p ÷ (fator efetivo × giro). Giro e fator são ilustrativos; o dado identifica só o produto p.",
]
for k, t in enumerate(notes):
    W(wsv, f"A{SV_LAST + 5 + k}", "• " + t)
for col, w in zip("ABCDEFGHIJKLMNOPQRST", [8, 40, 40, 9, 46, 9, 9, 9, 9, 8, 8, 8, 8, 8, 8, 8, 8, 8, 10, 60]):
    wsv.column_dimensions[col].width = w

# ---- Alocacao ----
wa = wb.create_sheet("Alocacao")
title(wa, "Alocação navio × serviço × período (inclui docagens e fretamento a terceiros)",
      "Fonte principal: tabelas de frota dos releases trimestrais (alocação no fim de cada trimestre) e eventos citados no texto.")
header_row(wa, 4, 1, ["Navio", "Serviço", "Início", "Fim", "Fundamento"])
for i, (n, s, a, b, why) in enumerate(STINTS):
    rr = 5 + i
    W(wa, f"A{rr}", n, fill=FILL_IN); W(wa, f"B{rr}", s, fill=FILL_IN, al=CEN)
    W(wa, f"C{rr}", a, fill=FILL_IN, nf=NFD); W(wa, f"D{rr}", b, fill=FILL_IN, nf=NFD); W(wa, f"E{rr}", why)
AL_FIRST, AL_LAST = 5, 4 + len(STINTS)
dv = DataValidation(type="list", formula1=f"=Servicos!$A${SV_FIRST}:$A${SV_LAST + len(OFFHIRE)}", allow_blank=False)
wa.add_data_validation(dv); dv.add(f"B{AL_FIRST}:B{AL_LAST + 10}")
for col, w in zip("ABCDE", [20, 9, 10, 10, 90]):
    wa.column_dimensions[col].width = w
wa.freeze_panes = "A5"

# ---- Calc_Navio ----
wn = wb.create_sheet("Calc_Navio")
title(wn, "Dias operando e slot-viagens por navio e trimestre",
      "Dias = sobreposição do período de alocação com o trimestre. Slot-viagens (mil) = capacidade × dias ÷ rotação ÷ 1.000.")
DC0 = 8                      # dias: H..
SC0 = DC0 + NQ + 1           # slot-viagens
dcol = lambda i: L(DC0 + i)
scol = lambda i: L(SC0 + i)
W(wn, f"{dcol(0)}3", "Dias operando no trimestre", F_S)
W(wn, f"{scol(0)}3", "Slot-viagens ofertadas (mil TEU)", F_S)
for lbl_row, lab in [(4, "Trimestre"), (5, "Início"), (6, "Fim"), (7, "Dias")]:
    W(wn, f"G{lbl_row}", lab, F_B)
for i in range(NQ):
    for blk in (dcol, scol):
        col = blk(i)
        if blk is dcol and i == 0:
            W(wn, f"{col}5", QS[0], F_B, FILL_IN, NFD)
        elif blk is dcol:
            W(wn, f"{col}5", f"=EDATE({dcol(i - 1)}5,3)", F_B, nf=NFD)
        else:
            W(wn, f"{col}5", f"={dcol(i)}5", F_B, nf=NFD)
        W(wn, f"{col}6", f"=EDATE({col}5,3)-1", nf=NFD)
        W(wn, f"{col}7", f"={col}6-{col}5+1", nf=NF0)
        W(wn, f"{col}4", f'=ROUNDUP(MONTH({col}5)/3,0)&"T"&RIGHT(YEAR({col}5),2)', F_W, FILL_H, al=CEN)
header_row(wn, 8, 1, ["Navio", "Serviço", "Início", "Fim", "Capacidade (TEU)", "Rotação (dias)"])
N_FIRST = 9
for i in range(len(STINTS)):
    rr = N_FIRST + i; ar = AL_FIRST + i
    W(wn, f"A{rr}", f"=Alocacao!A{ar}"); W(wn, f"B{rr}", f"=Alocacao!B{ar}", al=CEN)
    W(wn, f"C{rr}", f"=Alocacao!C{ar}", nf=NFD); W(wn, f"D{rr}", f"=Alocacao!D{ar}", nf=NFD)
    W(wn, f"E{rr}", f"=INDEX(Frota!$C$5:$C${FLEET_LAST},MATCH(A{rr},Frota!$A$5:$A${FLEET_LAST},0))", nf=NF0)
    W(wn, f"F{rr}", f"=IFERROR(INDEX(Servicos!$D${SV_FIRST}:$D${SV_LAST + len(OFFHIRE)},MATCH(B{rr},Servicos!$A${SV_FIRST}:$A${SV_LAST + len(OFFHIRE)},0)),0)", nf=NF0)
    for q in range(NQ):
        W(wn, f"{dcol(q)}{rr}", f"=MAX(0,MIN($D{rr},{dcol(q)}$6)-MAX($C{rr},{dcol(q)}$5)+1)", nf=NF0)
        W(wn, f"{scol(q)}{rr}", f"=IF($F{rr}>0,$E{rr}*{dcol(q)}{rr}/$F{rr}/1000,0)", nf=NF1)
N_LAST = N_FIRST + len(STINTS) - 1
rr = N_LAST + 1
W(wn, f"A{rr}", "Total", F_B)
for q in range(NQ):
    W(wn, f"{dcol(q)}{rr}", f"=SUM({dcol(q)}{N_FIRST}:{dcol(q)}{N_LAST})", F_B, nf=NF0, border=TOP)
    W(wn, f"{scol(q)}{rr}", f"=SUM({scol(q)}{N_FIRST}:{scol(q)}{N_LAST})", F_B, nf=NF1, border=TOP)
for col, w in zip("ABCDEFG", [20, 8, 9, 9, 10, 8, 10]):
    wn.column_dimensions[col].width = w
for q in range(NQ):
    wn.column_dimensions[dcol(q)].width = 9; wn.column_dimensions[scol(q)].width = 9
wn.freeze_panes = f"{dcol(0)}9"

# ---- Calc_Servico ----
wc = wb.create_sheet("Calc_Servico")
title(wc, "Oferta e volume modelado por serviço e trade (mil TEU)",
      "Volume = slot-viagens × p × mix (parâmetros da janela selecionada em Backtest!D3).")
QC0 = 3
cq = lambda i: L(QC0 + i)
W(wc, "A4", "mil TEU", F_B)
for i in range(NQ):
    W(wc, f"{cq(i)}4", f"=Calc_Navio!{dcol(i)}4", F_W, FILL_H, al=CEN)
# bloco 1: navios-equivalentes
r0 = 6
W(wc, f"A{r0}", "Navios-equivalentes em operação (dias-navio ÷ dias do trimestre)", F_S)
NE_ROW = {}
for k, code in enumerate(SCODES):
    rr = r0 + 1 + k; NE_ROW[code] = rr
    W(wc, f"A{rr}", code, F_B); W(wc, f"B{rr}", f"=VLOOKUP(A{rr},Servicos!$A$6:$B$14,2,FALSE)", F_I)
    for i in range(NQ):
        W(wc, f"{cq(i)}{rr}", f"=SUMIF(Calc_Navio!$B${N_FIRST}:$B${N_LAST},$A{rr},Calc_Navio!{dcol(i)}${N_FIRST}:{dcol(i)}${N_LAST})/Calc_Navio!{dcol(i)}$7", nf=NF2)
rr = r0 + 1 + len(SCODES)
W(wc, f"A{rr}", "Total", F_B)
for i in range(NQ):
    W(wc, f"{cq(i)}{rr}", f"=SUM({cq(i)}{r0 + 1}:{cq(i)}{rr - 1})", F_B, nf=NF2, border=TOP)
# bloco 2: oferta
r1 = rr + 2
W(wc, f"A{r1}", "Slot-viagens ofertadas (mil TEU)", F_S)
OF_ROW = {}
for k, code in enumerate(SCODES):
    rr = r1 + 1 + k; OF_ROW[code] = rr
    W(wc, f"A{rr}", code, F_B); W(wc, f"B{rr}", f"=VLOOKUP(A{rr},Servicos!$A$6:$B$14,2,FALSE)", F_I)
    for i in range(NQ):
        W(wc, f"{cq(i)}{rr}", f"=SUMIF(Calc_Navio!$B${N_FIRST}:$B${N_LAST},$A{rr},Calc_Navio!{scol(i)}${N_FIRST}:{scol(i)}${N_LAST})", nf=NF1)
rr = r1 + 1 + len(SCODES); OF_TOT = rr
W(wc, f"A{rr}", "Total", F_B)
for i in range(NQ):
    W(wc, f"{cq(i)}{rr}", f"=SUM({cq(i)}{r1 + 1}:{cq(i)}{rr - 1})", F_B, nf=NF1, border=TOP)
# bloco 3: volume por servico x trade
r2 = rr + 2
W(wc, f"A{r2}", "Volume modelado por serviço × trade (mil TEU)", F_S)
TRADES = [("Cabotagem", "P"), ("Mercosul", "Q"), ("Feeder", "R")]
VT_ROW = {}
rr = r2
for code in SCODES:
    for tname, mixcol in TRADES:
        rr += 1; VT_ROW[(code, tname)] = rr
        W(wc, f"A{rr}", code, F_B); W(wc, f"B{rr}", tname)
        for i in range(NQ):
            W(wc, f"{cq(i)}{rr}", f"={cq(i)}${OF_ROW[code]}*Servicos!$I${SV_ROW[code]}*Servicos!${mixcol}${SV_ROW[code]}", nf=NF1)
VT_FIRST, VT_LAST = r2 + 1, rr
# bloco 4: totais por trade
r3 = rr + 2
W(wc, f"A{r3}", "Volume modelado por trade (mil TEU)", F_S)
TR_ROW = {}
for k, (tname, _) in enumerate(TRADES):
    rr = r3 + 1 + k; TR_ROW[tname] = rr
    W(wc, f"A{rr}", tname, F_B)
    for i in range(NQ):
        W(wc, f"{cq(i)}{rr}", f"=SUMIF($B${VT_FIRST}:$B${VT_LAST},$A{rr},{cq(i)}${VT_FIRST}:{cq(i)}${VT_LAST})", nf=NF1)
rr = r3 + 4; TR_ROW["Total"] = rr
W(wc, f"A{rr}", "Total", F_B)
for i in range(NQ):
    W(wc, f"{cq(i)}{rr}", f"=SUM({cq(i)}{r3 + 1}:{cq(i)}{rr - 1})", F_B, nf=NF1, border=TOP)
# bloco 5: volume por servico
r4 = rr + 2
W(wc, f"A{r4}", "Volume modelado por serviço (mil TEU)", F_S)
VS_ROW = {}
for k, code in enumerate(SCODES):
    rr = r4 + 1 + k; VS_ROW[code] = rr
    W(wc, f"A{rr}", code, F_B)
    for i in range(NQ):
        W(wc, f"{cq(i)}{rr}", f"=SUMIF($A${VT_FIRST}:$A${VT_LAST},$A{rr},{cq(i)}${VT_FIRST}:{cq(i)}${VT_LAST})", nf=NF1)
wc.column_dimensions["A"].width = 10; wc.column_dimensions["B"].width = 44
for i in range(NQ):
    wc.column_dimensions[cq(i)].width = 8.5
wc.freeze_panes = "C5"

# ---- Backtest ----
wbt = wb.create_sheet("Backtest")
title(wbt, "Backtest: volume publicado × modelado (mil TEU)",
      "Laranja = trimestre fora da amostra de calibração. 2026: trades não divulgados, comparação só no total.")
W(wbt, "A3", "Janela de calibração (1 ou 2)", F_B); W(wbt, "D3", 1, F_B, FILL_IN, al=CEN)
W(wbt, "E3", '=IF(Janela=1,"1 = calibra 1T22–4T24 e testa 1T25–2T26 fora da amostra","2 = calibra 1T22–4T25 (amostra completa); teste só 1S26")', F_I)
dvj = DataValidation(type="list", formula1='"1,2"', allow_blank=False); wbt.add_data_validation(dvj); dvj.add("D3")
wb.defined_names["Janela"] = DefinedName("Janela", attr_text="Backtest!$D$3")
BC0 = 3
bq = lambda i: L(BC0 + i)
W(wbt, "A5", "mil TEU", F_B)
for i in range(NQ):
    W(wbt, f"{bq(i)}5", QL[i], F_W, FILL_H, al=CEN)
W(wbt, "A6", "Amostra", F_I)
for i in range(NQ):
    W(wbt, f"{bq(i)}6", f'=IF({i + 1}<=IF(Janela=1,12,16),"Calibração","Teste")', F_I, al=CEN)
blocks = [("Publicado", 8), ("Modelado", 14), ("Erro (modelado − publicado)", 20), ("Erro %", 26)]
TR4 = ["Cabotagem", "Mercosul", "Feeder", "Total"]
PUBR = {"Cabotagem": RC, "Mercosul": RM, "Feeder": RF, "Total": RT}
for name, br in blocks:
    W(wbt, f"A{br}", name, F_S)
    for k, t in enumerate(TR4):
        rr = br + 1 + k
        W(wbt, f"A{rr}", t, F_B if t == "Total" else F_N)
        for i in range(NQ):
            col = bq(i)
            if name == "Publicado":
                f = f"=Volumes_Publicados!{qcol(i)}{PUBR[t]}"; nf = NF1
            elif name == "Modelado":
                f = f"=Calc_Servico!{cq(i)}{TR_ROW[t]}"; nf = NF1
            elif name.startswith("Erro ("):
                f = f'=IF(ISNUMBER({col}{9 + k}),{col}{15 + k}-{col}{9 + k},"n.d.")'; nf = NF1
            else:
                f = f'=IF(ISNUMBER({col}{9 + k}),{col}{15 + k}/{col}{9 + k}-1,"n.d.")'; nf = NFP
            W(wbt, f"{col}{rr}", f, F_B if t == "Total" else F_N, nf=nf, al=CEN if False else None)
# |erro %| auxiliar
W(wbt, "A32", "|Erro %| (auxiliar)", F_S)
for k, t in enumerate(TR4):
    rr = 33 + k
    W(wbt, f"A{rr}", t, F_I)
    for i in range(NQ):
        W(wbt, f"{bq(i)}{rr}", f'=IF(ISNUMBER({bq(i)}{27 + k}),ABS({bq(i)}{27 + k}),"")', F_I, nf=NFP)
# eventos
W(wbt, "A38", "Eventos que explicam o resíduo", F_S)
for i in range(NQ):
    W(wbt, f"{bq(i)}39", EVENTS[i], F_I, al=WRAP)
wbt.row_dimensions[39].height = 105
# resumo
last = bq(NQ - 1)
W(wbt, "A42", "Resumo do erro (MAPE = média do |erro %|)", F_S)
header_row(wbt, 43, 1, ["Trade", "", "MAPE calibração", "MAPE teste", "Viés calibração", "Viés teste", "Nº tri teste"])
for k, t in enumerate(TR4):
    rr = 44 + k; er = 33 + k; ep = 27 + k
    W(wbt, f"A{rr}", t, F_B if t == "Total" else F_N)
    for j, (lab, fn) in enumerate([("Calibração", "abs"), ("Teste", "abs")]):
        W(wbt, f"{L(3 + j)}{rr}", f'=IFERROR(SUMIFS($C{er}:${last}{er},$C$6:${last}$6,"{lab}")/COUNTIFS($C$6:${last}$6,"{lab}",$C{er}:${last}{er},">=0"),"n.d.")', nf=NFP)
    for j, lab in enumerate(["Calibração", "Teste"]):
        W(wbt, f"{L(5 + j)}{rr}", f'=IFERROR(SUMIFS($C{ep}:${last}{ep},$C$6:${last}$6,"{lab}")/COUNTIFS($C$6:${last}$6,"{lab}",$C{er}:${last}{er},">=0"),"n.d.")', nf=NFP)
    W(wbt, f"G{rr}", f'=COUNTIFS($C$6:${last}$6,"Teste",$C{er}:${last}{er},">=0")', nf=NF0)
# anual
W(wbt, "A50", "Anual (mil TEU)", F_S)
header_row(wbt, 51, 1, ["", "", "2022", "2023", "2024", "2025", "1S26"])
for k, t in enumerate(TR4):
    for j, (lab, base) in enumerate([("Publicado", 9), ("Modelado", 15)]):
        rr = 52 + k * 3 + j
        W(wbt, f"A{rr}", t, F_B if t == "Total" else F_N); W(wbt, f"B{rr}", lab)
        for y, (yy, s0) in enumerate(YRS):
            n = 2 if yy == "1S26" else 4
            rng = f"{bq(s0)}{base + k}:{bq(s0 + n - 1)}{base + k}"
            W(wbt, f"{L(3 + y)}{rr}", f'=IF(COUNT({rng})={n},SUM({rng}),"n.d.")', nf=NF1)
    rr = 52 + k * 3 + 2
    W(wbt, f"B{rr}", "Erro %", F_I)
    for y in range(len(YRS)):
        cl = L(3 + y)
        W(wbt, f"{cl}{rr}", f'=IFERROR({cl}{rr - 1}/{cl}{rr - 2}-1,"n.d.")', F_I, nf=NFP)
wbt.column_dimensions["A"].width = 26; wbt.column_dimensions["B"].width = 10
for i in range(NQ):
    wbt.column_dimensions[bq(i)].width = 11
wbt.freeze_panes = "C6"
# graficos
ch = LineChart(); ch.title = "Total: publicado × modelado (mil TEU)"; ch.height = 7.5; ch.width = 22
for rr_, nm_ in [(12, "Publicado"), (18, "Modelado")]:
    ch.series.append(Series(Reference(wbt, min_col=BC0, max_col=BC0 + NQ - 1, min_row=rr_, max_row=rr_), title=nm_))
ch.set_categories(Reference(wbt, min_col=BC0, max_col=BC0 + NQ - 1, min_row=5, max_row=5))
ch.series[0].graphicalProperties.line.solidFill = "1F3864"
ch.series[1].graphicalProperties.line.solidFill = "9DC3A8"; ch.series[1].graphicalProperties.line.dashStyle = "dash"
wbt.add_chart(ch, "I42")
bc = BarChart(); bc.title = "Erro % do total por trimestre"; bc.height = 7.5; bc.width = 22
bc.add_data(Reference(wbt, min_col=BC0, max_col=BC0 + NQ - 1, min_row=30, max_row=30), from_rows=True, titles_from_data=False)
bc.set_categories(Reference(wbt, min_col=BC0, max_col=BC0 + NQ - 1, min_row=5, max_row=5))
bc.legend = None; bc.series[0].graphicalProperties.solidFill = "1F3864"
bc.y_axis.numFmt = "0%"; bc.series[0].invertIfNegative = False
for chx in (ch, bc):
    chx.x_axis.delete = False; chx.y_axis.delete = False
ch.y_axis.numFmt = "0"; ch.x_axis.tickLblPos = "low"; bc.x_axis.tickLblPos = "low"
for s_ in ch.series:
    s_.smooth = False
wbt.add_chart(bc, "I58")
# laranja nos trimestres fora da amostra
wbt.conditional_formatting.add(f"C6:{last}30", FormulaRule(formula=['C$6="Teste"'], fill=PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")))

# ---- Reconciliacao ----
wr = wb.create_sheet("Reconciliacao")
title(wr, "Reconciliação: produtividade implícita que fecha o volume publicado",
      "Fator de ajuste = publicado ÷ modelado por trade (2026: fator do total). Produtividade implícita do serviço = p × Σ(mix × fator).")
RC0 = 3
rq = lambda i: L(RC0 + i)
W(wr, "A4", "", F_B)
for i in range(NQ):
    W(wr, f"{rq(i)}4", QL[i], F_W, FILL_H, al=CEN)
W(wr, "A5", "Fator de ajuste (publicado ÷ modelado)", F_S)
FA_ROW = {}
for k, t in enumerate(["Cabotagem", "Mercosul", "Feeder"]):
    rr = 6 + k; FA_ROW[t] = rr
    W(wr, f"A{rr}", t)
    for i in range(NQ):
        b = bq(i)
        W(wr, f"{rq(i)}{rr}", f"=IF(ISNUMBER(Backtest!{b}{9 + k}),IFERROR(Backtest!{b}{9 + k}/Backtest!{b}{15 + k},1),Backtest!{b}12/Backtest!{b}18)", nf=NF2)
W(wr, "A9", "Total", F_B)
for i in range(NQ):
    W(wr, f"{rq(i)}9", f"=Backtest!{bq(i)}12/Backtest!{bq(i)}18", F_B, nf=NF2)
W(wr, "A11", "Produtividade implícita (TEU por slot-viagem)", F_S)
PI_ROW = {}
for k, code in enumerate(SCODES):
    rr = 12 + k; PI_ROW[code] = rr
    sr = SV_ROW[code]
    W(wr, f"A{rr}", code, F_B)
    W(wr, f"B{rr}", f"=Servicos!I{sr}", F_I, nf=NF2)
    for i in range(NQ):
        col = rq(i)
        W(wr, f"{col}{rr}", f"=IF(Calc_Servico!{cq(i)}{OF_ROW[code]}>0,Servicos!$I${sr}*(Servicos!$P${sr}*{col}$6+Servicos!$Q${sr}*{col}$7+Servicos!$R${sr}*{col}$8),\"\")", nf=NF2)
W(wr, "B11", "p calibrado", F_I)
r = 12 + len(SCODES) + 1
W(wr, f"A{r}", "Volume reconciliado por serviço (mil TEU)", F_S)
VR_ROW = {}
for k, code in enumerate(SCODES):
    rr = r + 1 + k; VR_ROW[code] = rr
    W(wr, f"A{rr}", code, F_B)
    for i in range(NQ):
        W(wr, f"{rq(i)}{rr}", f"=N({rq(i)}{PI_ROW[code]})*Calc_Servico!{cq(i)}{OF_ROW[code]}", nf=NF1)
rr = r + 1 + len(SCODES); VR_TOT = rr
W(wr, f"A{rr}", "Total reconciliado", F_B)
for i in range(NQ):
    W(wr, f"{rq(i)}{rr}", f"=SUM({rq(i)}{r + 1}:{rq(i)}{rr - 1})", F_B, nf=NF1, border=TOP)
W(wr, f"A{rr + 1}", "Diferença vs publicado", F_I)
for i in range(NQ):
    W(wr, f"{rq(i)}{rr + 1}", f"={rq(i)}{rr}-Backtest!{bq(i)}12", F_I, nf=NF1)
# KPIs
r = rr + 3
W(wr, f"A{r}", "Indicadores de produtividade da frota", F_S)
kpis = [
    ("Capacidade média em operação (TEU)", "cap"),
    ("Slot-viagens ofertadas (mil TEU)", "off"),
    ("TEU por slot-viagem (total)", "tps"),
    ("TEU por TEU de capacidade (anualizado)", "tpc"),
]
KPI_ROW = {}
for k, (lab, key) in enumerate(kpis):
    rr = r + 1 + k; KPI_ROW[key] = rr
    W(wr, f"A{rr}", lab)
    for i in range(NQ):
        col = rq(i); dc = dcol(i)
        if key == "cap":
            f = f"=SUMPRODUCT((Calc_Navio!$F${N_FIRST}:$F${N_LAST}>0)*Calc_Navio!$E${N_FIRST}:$E${N_LAST}*Calc_Navio!{dc}${N_FIRST}:{dc}${N_LAST})/Calc_Navio!{dc}$7"; nf = NF0
        elif key == "off":
            f = f"=Calc_Servico!{cq(i)}{OF_TOT}"; nf = NF1
        elif key == "tps":
            f = f"=Backtest!{bq(i)}12/{col}{r + 2}"; nf = NF2
        else:
            f = f"=Backtest!{bq(i)}12*1000*(365/Calc_Navio!{dc}$7)/{col}{r + 1}"; nf = NF1
        W(wr, f"{col}{rr}", f, nf=nf)
# ponte anual por servico
r = r + 6
W(wr, f"A{r}", "Ponte anual por serviço: Δ volume = efeito oferta + efeito produtividade (mil TEU)", F_S)
YA = [("2022", 0), ("2023", 4), ("2024", 8), ("2025", 12)]
hdr = ["Serviço"]
for (y0, _), (y1, _) in zip(YA[:-1], YA[1:]):
    hdr += [f"{y0}→{y1} oferta", f"{y0}→{y1} produtiv.", f"{y0}→{y1} Δ total"]
header_row(wr, r + 1, 1, hdr)
for k, code in enumerate(SCODES):
    rr = r + 2 + k
    W(wr, f"A{rr}", code, F_B)
    for j, ((y0, s0), (y1, s1)) in enumerate(zip(YA[:-1], YA[1:])):
        o0 = f"SUM(Calc_Servico!{cq(s0)}{OF_ROW[code]}:{cq(s0 + 3)}{OF_ROW[code]})"
        o1 = f"SUM(Calc_Servico!{cq(s1)}{OF_ROW[code]}:{cq(s1 + 3)}{OF_ROW[code]})"
        v0 = f"SUM({rq(s0)}{VR_ROW[code]}:{rq(s0 + 3)}{VR_ROW[code]})"
        v1 = f"SUM({rq(s1)}{VR_ROW[code]}:{rq(s1 + 3)}{VR_ROW[code]})"
        p0 = f"IF({o0}>0,{v0}/{o0},IF({o1}>0,{v1}/{o1},0))"  # servico novo: todo o delta e efeito oferta
        c1, c2, c3 = L(2 + 3 * j), L(3 + 3 * j), L(4 + 3 * j)
        W(wr, f"{c1}{rr}", f"=({o1}-{o0})*{p0}", nf=NF1)
        W(wr, f"{c3}{rr}", f"={v1}-{v0}", nf=NF1)
        W(wr, f"{c2}{rr}", f"={c3}{rr}-{c1}{rr}", nf=NF1)
rr = r + 2 + len(SCODES)
W(wr, f"A{rr}", "Total", F_B)
for j in range(9):
    cl = L(2 + j)
    W(wr, f"{cl}{rr}", f"=SUM({cl}{r + 2}:{cl}{rr - 1})", F_B, nf=NF1, border=TOP)
wr.column_dimensions["A"].width = 40; wr.column_dimensions["B"].width = 11
for i in range(NQ):
    wr.column_dimensions[rq(i)].width = 11
wr.freeze_panes = "C5"
GR = rr + 3
W(wr, f"A{GR}", "Dados do gráfico: produtividade implícita por grupo de serviço (#N/A = serviço inexistente)", F_S)
GROUPS = [("SAS + SAM", ["SAS", "SAM"]), ("SEA (antigo + reformulado)", ["SEA", "SEA25"]), ("Shuttles (SSV + SSR + SSN)", ["SSV", "SSR", "SSN"])]
for k, (gname, members) in enumerate(GROUPS):
    rg = GR + 1 + k
    W(wr, f"A{rg}", gname, F_I)
    for i in range(NQ):
        vol = "+".join(f"{rq(i)}{VR_ROW[m]}" for m in members)
        off = "+".join(f"Calc_Servico!{cq(i)}{OF_ROW[m]}" for m in members)
        W(wr, f"{rq(i)}{rg}", f"=IF(({off})>0,({vol})/({off}),NA())", F_I, nf=NF2)
wr.conditional_formatting.add(f"C{GR + 1}:{rq(NQ - 1)}{GR + 3}", FormulaRule(formula=[f"ISNA(C{GR + 1})"], font=Font(name=ARIAL, size=9, color="BFBFBF")))
lc = LineChart(); lc.title = "Produtividade implícita por grupo de serviço (TEU por slot-viagem)"; lc.height = 8; lc.width = 24
for k, (gname, _) in enumerate(GROUPS):
    rg = GR + 1 + k
    lc.series.append(Series(Reference(wr, min_col=RC0, max_col=RC0 + NQ - 1, min_row=rg, max_row=rg), title=gname))
lc.set_categories(Reference(wr, min_col=RC0, max_col=RC0 + NQ - 1, min_row=4, max_row=4))
lc.x_axis.delete = False; lc.y_axis.delete = False; lc.y_axis.numFmt = "0.0"; lc.x_axis.tickLblPos = "low"
lc.display_blanks = "gap"
for s_ in lc.series:
    s_.smooth = False
for s_, col_ in zip(lc.series, ["1F3864", "9DC3A8", "C55A11"]):
    s_.graphicalProperties.line.solidFill = col_; s_.graphicalProperties.line.width = 22000
wr.add_chart(lc, f"A{GR + 5}")

# ---- Checagens ----
wk = wb.create_sheet("Checagens")
title(wk, "Checagens de consistência", "Capacidade de frota reportada × modelo e fechamento dos números anuais publicados.")
KC0 = 3
kq = lambda i: L(KC0 + i)
for i in range(NQ):
    W(wk, f"{kq(i)}4", QL[i], F_W, FILL_H, al=CEN)
W(wk, "A5", "Capacidade nominal reportada (TEU)", F_B)
W(wk, "A6", "Modelo: navios na frota no fim do tri (exceto docagem)")
W(wk, "A7", "Modelo: capacidade média em operação + fretado")
W(wk, "A8", "Diferença (fim do tri − reportado)", F_I)
W(wk, "A9", "Diferença (média − reportado)", F_I)
for i in range(NQ):
    col = kq(i); dc = dcol(i)
    W(wk, f"{col}5", CAP_REP[i], fill=FILL_IN, nf=NF0)
    W(wk, f"{col}6", f'=SUMPRODUCT((Calc_Navio!$C${N_FIRST}:$C${N_LAST}<=Calc_Navio!{dc}$6)*(Calc_Navio!$D${N_FIRST}:$D${N_LAST}>=Calc_Navio!{dc}$6)*(Calc_Navio!$B${N_FIRST}:$B${N_LAST}<>"DOC")*Calc_Navio!$E${N_FIRST}:$E${N_LAST})', nf=NF0)
    W(wk, f"{col}7", f'=SUMPRODUCT((Calc_Navio!$B${N_FIRST}:$B${N_LAST}<>"DOC")*Calc_Navio!$E${N_FIRST}:$E${N_LAST}*Calc_Navio!{dc}${N_FIRST}:{dc}${N_LAST})/Calc_Navio!{dc}$7', nf=NF0)
    W(wk, f"{col}8", f"={col}6-{col}5", F_I, nf=NF0)
    W(wk, f"{col}9", f"={col}7-{col}5", F_I, nf=NF0)
W(wk, "A10", "Leitura: até 2025 o release reporta a frota no fim do trimestre (inclui o Discovery fretado em 2022); em 2026 o número bate com a média do período (docagens do 2T26).", F_I)
W(wk, "A11", "Diferenças pequenas: Belmonte 3.500 × 3.534 (2T24) e Evolution 3.150 × 3.158 (1T24) nas tabelas de frota.", F_I)
W(wk, "A12", "3T22: a Log-In manteve o Jacarandá (em docagem) na capacidade reportada; no 1T24 excluiu o Discovery em docagem. O critério do release muda; o modelo exclui docagem sempre.", F_I)
W(wk, "A14", "Fechamento dos números anuais publicados", F_S)
header_row(wk, 15, 1, ["Item", "", "Publicado", "Série reconstruída", "Diferença"])
Y = {y: s0 for y, s0 in YRS}
def ysum(row, y):
    s0 = Y[y]; return f"SUM(Volumes_Publicados!{qcol(s0)}{row}:{qcol(s0 + 3)}{row})"
checks = [
    ("Total 2022", f"={FR('T22A')}", f"={ysum(RT, '2022')}", NF1),
    ("Total 2023", f"={FR('T23A')}", f"={ysum(RT, '2023')}", NF1),
    ("Total 2024", f"={FR('T24A')}", f"={ysum(RT, '2024')}", NF1),
    ("Total 2025", f"={FR('T25A')}", f"={ysum(RT, '2025')}", NF1),
    ("Cabotagem 2024 (var. a/a)", f"={FR('g_C24A')}", f"={ysum(RC, '2024')}/{ysum(RC, '2023')}-1", NFP),
    ("Cabotagem 2025", f"={FR('C25A')}", f"={ysum(RC, '2025')}", NF1),
    ("Feeder 2024", f"={FR('F24A')}", f"={ysum(RF, '2024')}", NF1),
    ("Feeder 2025", f"={FR('F25A')}", f"={ysum(RF, '2025')}", NF1),
    ("Feeder 3T23 (via 3T24 ÷ 1,89)", f"={FR('F3T23')}", f"={FR('F3T24')}/(1+{FR('g_F3T24')})", NF1),
    ("Mercosul 1T23 (var. a/a)", f"={FR('g_M1T23')}", f"=Volumes_Publicados!{qcol(4)}{RM}/Volumes_Publicados!{qcol(0)}{RM}-1", NFP),
    ("Mercosul 3T24 (var. a/a)", f"={FR('g_M3T24')}", f"=Volumes_Publicados!{qcol(10)}{RM}/Volumes_Publicados!{qcol(6)}{RM}-1", NFP),
    ("Mercosul 4T24 (var. a/a)", f"={FR('g_M4T24')}", f"=Volumes_Publicados!{qcol(11)}{RM}/Volumes_Publicados!{qcol(7)}{RM}-1", NFP),
    ("Mercosul 2024 (var. a/a)", f"={FR('g_M24A')}", f"={ysum(RM, '2024')}/{ysum(RM, '2023')}-1", NFP),
    ("Mercosul 1T25 (var. a/a)", f"={FR('g_M1T25')}", f"=Volumes_Publicados!{qcol(12)}{RM}/Volumes_Publicados!{qcol(8)}{RM}-1", NFP),
    ("Mercosul 2T25 (var. a/a)", f"={FR('g_M2T25')}", f"=Volumes_Publicados!{qcol(13)}{RM}/Volumes_Publicados!{qcol(9)}{RM}-1", NFP),
    ("Mercosul 3T25 (var. a/a)", f"={FR('g_M3T25')}", f"=Volumes_Publicados!{qcol(14)}{RM}/Volumes_Publicados!{qcol(10)}{RM}-1", NFP),
    ("Mercosul 4T25 (var. a/a)", f"={FR('g_M4T25')}", f"=Volumes_Publicados!{qcol(15)}{RM}/Volumes_Publicados!{qcol(11)}{RM}-1", NFP),
    ("Mercosul 2025 (var. a/a)", f"={FR('g_M25A')}", f"={ysum(RM, '2025')}/{ysum(RM, '2024')}-1", NFP),
]
for k, (lab, fp, fs, nf) in enumerate(checks):
    rr = 16 + k
    W(wk, f"A{rr}", lab); W(wk, f"C{rr}", fp, nf=nf); W(wk, f"D{rr}", fs, nf=nf)
    W(wk, f"E{rr}", f"=D{rr}-C{rr}", F_I, nf=nf)
W(wk, f"A{17 + len(checks)}", "Mercosul trimestral de 2023–25 é resíduo; as variações a/a recalculadas diferem poucos p.p. das publicadas por arredondamento dos volumes de 1 casa.", F_I)
wk.column_dimensions["A"].width = 52; wk.column_dimensions["B"].width = 2
for i in range(NQ):
    wk.column_dimensions[kq(i)].width = 10
wk.column_dimensions["D"].width = 16

# ---- Evidencia_Rotacao ----
we = wb.create_sheet("Evidencia_Rotacao")
title(we, "Evidência de rotação: programação de navios da Log-In (atualizada em 30/09/2026)",
      "Fonte: loginlogistica.com.br/programacao/programacao-de-navios (PDF da planilha completa). Viagem = nº sequencial do navio.")
header_row(we, 4, 1, ["Serviço", "Navios na programação", "Rotação de portos", "Observação de volta", "Dias por volta", "Navios × frequência"])
EVR = [
    ("SAS", "Log-In Resiliente (498/499), Mercosul Suape (359/360, parceiro), Log-In Experience (29/30), Log-In Jatobá (287/288)",
     "BUE → La Plata → Rio Grande → Navegantes → Santos → Suape → Pecém → Salvador → Santos → Navegantes → BUE",
     "Resiliente v.498 ETA BUE 30/08 → v.499 ETA BUE 04/10", 35, "5 posições × 7 dias (uma vaga sem navio na janela)"),
    ("SEA", "Log-In Jacarandá (197/198), Experience (29, phase-out), Endurance (84/85), Evolution (55/56), Polaris (133, phase-in)",
     "Santos → Navegantes → Salvador → Suape → Pecém → Itacoatiara → Manaus → Suape → Vitória → Santos",
     "Jacarandá v.197 ETA SSZ 30/08 → v.198 27/09; Endurance 14/09 → 11/10; Evolution 20/09 → 17/10", 28, "4 navios × 7 dias"),
    ("SSR", "Log-In Pantanal (614–621)", "Santos → Rio (ICTSI) → Rio (Multi) → Vitória → Santos",
     "Pantanal v.614 ETA SSZ 02/09 → v.615 09/09 → v.616 16/09", 7, "1 navio, semanal"),
    ("SSV", "Log-In Discovery (80–87)", "Rio (Multi) → Vitória → Rio (Multi)",
     "Discovery v.82 ETA RIO 18/09 → v.83 25/09 → v.84 03/10", 7, "1 navio, semanal"),
    ("SMN", "Pedro Álvares, F. Magalhães, Am. Vespucio, Sebastião Caboto (navios de terceiros)",
     "Itaguaí → Santos → Salvador → Pecém → Manaus", "Serviço com navios de parceiro (compra/troca de espaço)", None, "fora do modelo (sem navio Log-In)"),
    ("SPA", "Felixstowe (terceiro)", "Buenos Aires → Santos", "Serviço novo com navio de terceiro", None, "fora do modelo"),
]
for k, row in enumerate(EVR):
    rr = 5 + k
    for j, v in enumerate(row):
        W(we, f"{L(1 + j)}{rr}", v, fill=(FILL_IN if j == 4 and v else None), al=WRAP, nf=(NF0 if j == 4 else None))
W(we, "A12", "Checagem pela numeração de viagem: o Experience (em operação desde ~ago/24) está na viagem 29–30 em set/26 → ~26 meses ÷ 29 voltas ≈ 27 dias por volta, coerente com rotações de 28–35 dias.", F_I)
W(we, "A13", "Em set/26 a alocação já difere da do 4T25 (Resiliente no SAS, Discovery no SSV, Endurance no SEA): o modelo usa as tabelas trimestrais até 4T25 e as mantém em 1S26.", F_I)
for col, w in zip("ABCDEF", [8, 55, 60, 55, 10, 32]):
    we.column_dimensions[col].width = w

# ---- ANTAQ_Plug ----
wq = wb.create_sheet("ANTAQ_Plug")
title(wq, "Layout para plugar o Estatístico Aquaviário da ANTAQ (por navio e escala)",
      "Painel/downloads da ANTAQ fora do ar em 01/10/2026 (aviso de manutenção). Quando voltar: baixar <ano>Atracacao.txt e <ano>Carga.txt.")
steps = [
    "1) Preencher o IMO dos navios na aba Frota.",
    "2) Filtrar Atracacao.txt pelos IMOs → uma linha por escala (porto, data de atracação/desatracação).",
    "3) Juntar Carga.txt por IDAtracacao → TEU embarcado/desembarcado, cheio/vazio, tipo de navegação (cabotagem/longo curso), origem/destino.",
    "4) Substituir premissas: rotação (dias entre escalas sucessivas no mesmo porto), dias de docagem (lacunas sem escala) e serviço (sequência de portos).",
    "5) Volume por navio = Σ TEU embarcado (cheio + vazio, se a Log-In contar vazios) → validar contra a série publicada e quebrar Feeder × Cabotagem × Mercosul por navio.",
]
for k, s in enumerate(steps):
    W(wq, f"A{4 + k}", s)
header_row(wq, 10, 1, ["IDAtracacao", "IMO", "Navio", "Porto", "Data atracação", "Data desatracação", "Tipo navegação",
                       "Sentido", "Cheio/Vazio", "TEU", "Origem", "Destino", "Serviço (inferido)", "Trade (inferido)"])
for rr in range(11, 16):
    for j in range(14):
        W(wq, f"{L(1 + j)}{rr}", None, fill=FILL_IN)
wq.column_dimensions["A"].width = 14
for j in range(2, 15):
    wq.column_dimensions[L(j)].width = 14

# ---- Fontes ----
wo = wb.create_sheet("Fontes")
title(wo, "Fontes")
SRC = [
    ("Planilha de Resultados (RI) – aba Volumes", "https://ri.loginlogistica.com.br/informacoes-aos-investidores/planilhas-suporte-de-resultados/"),
    ("Central de Resultados (releases 1T22–2T26)", "https://ri.loginlogistica.com.br/informacoes-aos-investidores/central-de-resultados/"),
    ("Programação de navios (PDF de 30/09/2026)", "https://www.loginlogistica.com.br/wp-content/uploads/ship_schedule/programacao-de-navio-20260930154506.pdf"),
    ("ABAC – números do setor (mercado por segmento)", "https://abac-br.org.br/cabotagem/numeros-do-setor/"),
    ("ANTAQ – Estatístico Aquaviário (indisponível em out/26)", "https://estatistica.antaq.gov.br/ea/sense/download.html"),
]
for k, (a, b) in enumerate(SRC):
    W(wo, f"A{4 + k}", a); W(wo, f"B{4 + k}", b)
W(wo, "A11", "Releases usados: 1T22, 2T22, 3T22, 4T22, 1T23, 2T23, 3T23, 4T23, 1T24, 2T24, 3T24, 4T24, 1T25, 2T25, 3T25, 4T25, 1T26, 2T26 (seções Volumes, destaques, Investimentos e tabela de frota).", F_I)
wo.column_dimensions["A"].width = 52; wo.column_dimensions["B"].width = 110

# ordem das abas
order = ["Leia-me", "Backtest", "Reconciliacao", "Volumes_Publicados", "Servicos", "Alocacao", "Frota",
         "Calc_Navio", "Calc_Servico", "Checagens", "Evidencia_Rotacao", "ANTAQ_Plug", "Fontes"]
wb._sheets = [wb[n] for n in order]
wb.active = 1
wb.save(OUT)

# ----------------------------------------------------------------------------------------------------------
# 5. LOG DA CALIBRACAO (para conferir com o Excel recalculado)
# ----------------------------------------------------------------------------------------------------------
def _finalizar_uma_vez(path):
    import os, signal
    import win32com.client, win32process
    xl = win32com.client.DispatchEx("Excel.Application")
    pid = win32process.GetWindowThreadProcessId(xl.Hwnd)[1]
    xl.Visible = False; xl.DisplayAlerts = False
    erros = []
    try:
        wbx = xl.Workbooks.Open(str(path))
        xl.CalculateFull()
        for wsx in wbx.Worksheets:
            wsx.Activate(); xl.ActiveWindow.GridlineColor = 0xF2F2F2; xl.ActiveWindow.Zoom = 90
            vals = wsx.UsedRange.Value
            vals = vals if isinstance(vals, tuple) else ((vals,),)
            for i_, row_ in enumerate(vals):
                for j_, v_ in enumerate(row_ if isinstance(row_, tuple) else (row_,)):
                    if isinstance(v_, int) and -2146826288 <= v_ <= -2146826246 and v_ != -2146826246:
                        erros.append(f"{wsx.Name}!R{wsx.UsedRange.Row + i_}C{wsx.UsedRange.Column + j_}")
        wbx.Worksheets("Backtest").Activate()
        wbx.Save(); wbx.Close(False)
        xl.Quit()
    finally:
        del xl
        try:
            os.kill(pid, signal.SIGTERM)  # garante que a instancia oculta nao fique pendurada
        except OSError:
            pass
    return erros


def finalizar_no_excel(path, tentativas=3):
    """Recalcula no Excel (COM), aplica grade cinza #F2F2F2 e devolve erros de formula reais."""
    for k in range(tentativas):
        try:
            return _finalizar_uma_vez(path)
        except Exception as e:  # COM instavel: tenta de novo com instancia nova
            ultimo = e
    raise ultimo


if __name__ == "__main__":
    print("erros de formula:", finalizar_no_excel(OUT) or "nenhum")
    np.set_printoptions(linewidth=200)
    print("serie cab", np.round(CAB[:16], 2)); print("serie mer", np.round(MER[:16], 2)); print("serie fee", np.round(FEE[:16], 2))
    for nm, par, pred in [("janela 1", PAR_A, PRED_A), ("janela 2", PAR_B, PRED_B)]:
        print(nm, {k: (round(v[0], 3), tuple(round(x, 2) for x in v[1])) for k, v in par.items()})
        print("  pred total", np.round(pred, 1))
        print("  err %     ", np.round((pred / np.array(TOTAL) - 1) * 100, 1))
    print("salvo:", OUT)
